import db as dbmod


def _client(tmp_path, monkeypatch, name="s.db"):
    monkeypatch.setattr(dbmod, "DB_PATH", str(tmp_path / name))
    dbmod.init_db()
    import app as appmod
    appmod.app.testing = True
    return appmod.app.test_client()


def _seed_products():
    conn = dbmod.get_db()
    conn.execute("""INSERT INTO products (raw, clean, category, barcode, brand, unit, weight, notes, grocy_track)
                    VALUES ('MILK 1L','Milk','Dairy','0123','MEV','L','1L','fresh',1)""")
    conn.execute("""INSERT INTO products (raw, clean, category, barcode)
                    VALUES ('BREAD','Bread','Bakery','456')""")
    conn.commit()
    conn.close()


def _workbook(c):
    import io
    r = c.get("/export-products.xlsx")
    assert r.status_code == 200
    assert "spreadsheet" in r.headers.get("Content-Type", "")
    import pandas as pd
    return pd.read_excel(io.BytesIO(r.get_data()), sheet_name=None)


def test_export_products_columns_and_text(tmp_path, monkeypatch):
    c = _client(tmp_path, monkeypatch)
    _seed_products()
    wb = _workbook(c)
    assert set(wb) == {"Products"}
    df = wb["Products"]
    assert len(df) == 2
    for col in ["raw", "clean", "category", "barcode", "brand", "unit",
                "weight", "notes", "grocy_track"]:
        assert col in df.columns, col
    assert "grocy_id" not in df.columns and "grocy_name" not in df.columns
    row = df[df["raw"] == "MILK 1L"].iloc[0]
    assert row["clean"] == "Milk"
    # raw xlsx cells stay text (pd.read_excel infers numbers on read)
    import io as _io
    import openpyxl
    r = c.get("/export-products.xlsx")
    ws = openpyxl.load_workbook(_io.BytesIO(r.get_data()))["Products"]
    hdr = [cell.value for cell in ws[1]]
    bi = hdr.index("barcode")
    vals = {ws.cell(row=n, column=1).value: (ws.cell(row=n, column=bi + 1).value,
            ws.cell(row=n, column=bi + 1).data_type) for n in range(2, ws.max_row + 1)}
    assert vals["MILK 1L"] == ("0123", "s")  # leading zero survives as text


def _upload_xlsx(c, rows, filename="friend.xlsx"):
    import io as _io
    import pandas as pd
    buf = _io.BytesIO()
    with pd.ExcelWriter(buf) as w:
        pd.DataFrame(rows).to_excel(w, sheet_name="Products", index=False)
    buf.seek(0)
    return c.post("/import-products", data={"file": (buf, filename)},
                  content_type="multipart/form-data", follow_redirects=True)


def test_import_clean_rows_auto_added(tmp_path, monkeypatch):
    c = _client(tmp_path, monkeypatch, "i.db")
    _seed_products()
    r = _upload_xlsx(c, [
        {"raw": "MILK 1L", "clean": "Milk", "category": "Dairy", "barcode": "0123",
         "brand": "MEV", "unit": "L", "weight": "1L", "notes": "fresh", "grocy_track": 1},
        {"raw": "CHEESE", "clean": "Cheese", "category": "Dairy", "barcode": "789"},
    ])
    assert r.status_code == 200
    conn = dbmod.get_db()
    row = conn.execute("SELECT clean, barcode FROM products WHERE raw='CHEESE'").fetchone()
    assert row is not None and tuple(row) == ("Cheese", "789")
    n = conn.execute("SELECT COUNT(*) FROM products").fetchone()[0]
    conn.close()
    assert n == 3
    assert "0 conflicts" in r.get_data(as_text=True)


def test_import_detects_conflicts(tmp_path, monkeypatch):
    c = _client(tmp_path, monkeypatch, "i2.db")
    _seed_products()
    r = _upload_xlsx(c, [
        {"raw": "MILK 1L", "clean": "Milk Friend", "category": "Dairy", "barcode": "0123"},
    ])
    body = r.get_data(as_text=True)
    assert "Milk Friend" in body and "Milk" in body
    assert "keep-new" in body and "keep-local" in body and "keep-both" in body


def test_resolve_keep_new_and_keep_local(tmp_path, monkeypatch):
    c = _client(tmp_path, monkeypatch, "i3.db")
    _seed_products()
    _upload_xlsx(c, [
        {"raw": "MILK 1L", "clean": "Milk Friend", "category": "Dairy"},
        {"raw": "BREAD", "clean": "Bread Friend", "category": "Bakery"},
    ])
    r = c.post("/import-products/resolve",
               data={"choice:MILK 1L": "new", "choice:BREAD": "local"},
               follow_redirects=True)
    assert r.status_code == 200
    conn = dbmod.get_db()
    assert conn.execute("SELECT clean FROM products WHERE raw='MILK 1L'").fetchone()[0] == "Milk Friend"
    assert conn.execute("SELECT clean FROM products WHERE raw='BREAD'").fetchone()[0] == "Bread"
    assert conn.execute("SELECT COUNT(*) FROM products").fetchone()[0] == 2
    conn.close()


def test_resolve_keep_both_renames(tmp_path, monkeypatch):
    c = _client(tmp_path, monkeypatch, "i4.db")
    _seed_products()
    _upload_xlsx(c, [{"raw": "MILK 1L", "clean": "Milk Friend", "category": "Dairy"}],
                 filename="anna.xlsx")
    c.post("/import-products/resolve", data={"choice:MILK 1L": "both"},
           follow_redirects=True)
    conn = dbmod.get_db()
    rows = conn.execute("SELECT raw, clean FROM products ORDER BY raw").fetchall()
    conn.close()
    assert ("MILK 1L", "Milk") in [(r[0], r[1]) for r in rows]
    assert ("MILK 1L [from anna.xlsx]", "Milk Friend") in [(r[0], r[1]) for r in rows]


def test_barcode_collision_is_warning_only(tmp_path, monkeypatch):
    c = _client(tmp_path, monkeypatch, "i5.db")
    _seed_products()
    r = _upload_xlsx(c, [{"raw": "MILK OTHER", "clean": "Other Milk", "barcode": "0123"}])
    body = r.get_data(as_text=True)
    assert "0123" in body and "warning" in body.lower()
    conn = dbmod.get_db()
    assert conn.execute("SELECT COUNT(*) FROM products").fetchone()[0] == 3
    conn.close()


def test_backup_db_is_valid_sqlite(tmp_path, monkeypatch):
    import sqlite3
    c = _client(tmp_path, monkeypatch, "i6.db")
    _seed_products()
    r = c.get("/backup-db")
    assert r.status_code == 200
    assert r.headers.get("Content-Type", "") == "application/x-sqlite3"
    path = str(tmp_path / "bak.db")
    with open(path, "wb") as f:
        f.write(r.get_data())
    conn = sqlite3.connect(path)
    try:
        n = conn.execute("SELECT COUNT(*) FROM products").fetchone()[0]
    finally:
        conn.close()
    assert n == 2
